# --- 檔案: src/tools/report_generator_docx.py ---
import os
import logging
import json
import sqlite3
from pathlib import Path
from typing import List

import sys
SRC_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(SRC_DIR))
from db import database

log = logging.getLogger(__name__)

def create_docx_report(task_ids: List[int], db_conn: sqlite3.Connection) -> str:
    from docx import Document
    from docx.shared import Inches
    log.info(f"開始為任務 {task_ids} 生成 Word 報告...")
    document = Document()
    document.add_heading('綜合績效分析報告', level=1)
    for task_id in task_ids:
        task_data = database.get_analysis_task(db_conn, task_id=task_id)
        if not task_data:
            log.warning(f"找不到任務 {task_id} 的資料，已跳過。")
            document.add_paragraph(f"找不到任務 ID: {task_id} 的分析資料。")
            continue
        document.add_heading(f"分析報告: {task_data.get('filename', '未知檔案')}", level=2)
        json_path_str = task_data.get("stage1_json_path")
        if not json_path_str or not Path(json_path_str).exists():
            log.warning(f"任務 {task_id} 的 JSON 檔案路徑不存在或遺失。")
            document.add_paragraph("錯誤：找不到分析結果的 JSON 檔案。")
            continue
        with open(json_path_str, "r", encoding="utf-8") as f:
            report_data = json.load(f)
        add_report_content_to_document(document, report_data)
        chart_path = generate_placeholder_chart()
        document.add_picture(chart_path, width=Inches(6.0))
        os.remove(chart_path)
        document.add_page_break()
    REPORTS_DIR = Path(__file__).resolve().parent.parent.parent / "reports"
    REPORTS_DIR.mkdir(exist_ok=True)
    output_filename = f"generated_report_{task_ids[0]}.docx"
    output_path = REPORTS_DIR / output_filename
    document.save(output_path)
    log.info(f"Word 報告已成功儲存至: {output_path}")
    return str(output_path)

def add_report_content_to_document(document, data: dict):
    table = document.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = '項目'
    hdr_cells[1].text = '內容'
    summary_data = {
        "分析標的": data.get("symbol", "N/A"),
        "公司名稱": data.get("company_name", "N/A"),
        "報告類型": data.get("report_type", "N/A"),
        "分析師": data.get("analyst", "N/A"),
        "目標價": str(data.get("target_price", "N/A")),
        "評級": data.get("rating", "N/A"),
    }
    for key, value in summary_data.items():
        row_cells = table.add_row().cells
        row_cells[0].text = key
        row_cells[1].text = value
    document.add_heading('詳細分析摘要', level=3)
    document.add_paragraph(data.get("summary", "沒有提供摘要。"))
    performance = data.get("performance_analysis", {})
    if performance:
        document.add_heading('量化績效分析', level=3)
        perf_table = document.add_table(rows=1, cols=2)
        perf_table.style = 'Table Grid'
        perf_hdr_cells = perf_table.rows[0].cells
        perf_hdr_cells[0].text = '績效指標'
        perf_hdr_cells[1].text = '數值'
        for key, value in performance.items():
            if key != 'chart_data':
                row_cells = perf_table.add_row().cells
                row_cells[0].text = str(key)
                row_cells[1].text = str(value)

def generate_placeholder_chart() -> str:
    import plotly.graph_objects as go
    fig = go.Figure(data=go.Scatter(x=[1, 2, 3, 4], y=[10, 11, 12, 13], mode='markers+lines'))
    fig.update_layout(title_text="範例績效圖表")
    chart_path = Path(__file__).resolve().parent / "temp_chart.png"
    fig.write_image(chart_path, engine="kaleido")
    return str(chart_path)
